"""非结构化文档服务：解析 .md/.txt/.pdf/.docx，导入为 ``doc:<id>`` 数据源并支持检索问答。

轻量 RAG。Markdown 按 ``=== <id> ===`` 分块，
表格块 → pandas DataFrame → SQLite 表（走既有 text2sql 回答聚合/列问题），
文本块 → 重叠分块 → Milvus 向量检索后交给 LLM 回答。

存储于 ``data/docs.db``：
- ``doc_meta``：文档元信息（doc_id / filename / title / kind / created_at）
- ``doc_blocks``：分块记录（表格块存列名+样本行、文本块存分块文本）
- 每张提取出的表格另建一张 ``doc_<doc_id>_<i>`` 表
"""
import datetime
import io
import json
import logging
import re
import uuid
from pathlib import Path

import pandas as pd
from fastapi import UploadFile
from sqlalchemy import inspect, text

from app.config import settings
from app.database.engine import _engine_kwargs, _get_or_create_engine, invalidate_engine
from app.database.milvus import list_field_values, vector_delete, vector_index, vector_search
from app.errors import AppError
from app.services.datasource_service import (
    _clean_identifier,
    _sanitize_columns,
    read_upload_limited,
)
from app.services.chunking import chunk_text

logger = logging.getLogger(__name__)

_ALLOWED_DOC_EXTS = {".md", ".txt", ".pdf", ".docx"}

# Markdown 分块标记：形如「=== 1 ===」或「=== table_3 ===」
_BLOCK_MARKER = re.compile(r"^===\s*(.*?)\s*===\s*$", re.M)
_HTML_TABLE_RE = re.compile(r"<table\b.*?</table>", flags=re.S | re.I)


def _now_iso() -> str:
    return datetime.datetime.now(datetime.timezone.utc).isoformat()


# doc_id 由 uuid4().hex[:12] 生成，形状固定；任何不符合的值都不得进入 Milvus 过滤表达式
_DOC_ID_RE = re.compile(r"^[0-9a-f]{12}$")


def _safe_doc_id(doc_id: str) -> str:
    """校验 doc_id 形状后返回。

    doc_id 会被拼进 Milvus 过滤表达式（``doc_id == "<id>"``），裸拼时
    ``doc:x" or doc_id != "`` 形态可跨文档读写整个集合，故先按生成规则严格校验。
    """
    if not _DOC_ID_RE.fullmatch(str(doc_id or "")):
        raise AppError(f"无效的文档标识：{doc_id}", "DATASOURCE_UNAVAILABLE")
    return doc_id


def _extract_text(filename: str, content: bytes) -> str:
    """按扩展名提取纯文本；缺可选解析库时抛友好 AppError。"""
    ext = Path(filename).suffix.lower()
    if ext in (".md", ".txt"):
        return content.decode("utf-8", errors="ignore")
    if ext == ".docx":
        try:
            import docx  # type: ignore
        except Exception:
            raise AppError("解析 .docx 需要安装 python-docx，请先安装后重试。", "SYSTEM_ERROR")
        doc = docx.Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            rows = [
                "| " + " | ".join(cell.text for cell in row.cells) + " |"
                for row in table.rows
            ]
            parts.append("\n".join(rows))
        return "\n\n".join(parts)
    if ext == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore
        except Exception:
            raise AppError("解析 .pdf 需要安装 pypdf，请先安装后重试。", "SYSTEM_ERROR")
        reader = PdfReader(io.BytesIO(content))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    raise AppError("仅支持 Markdown / TXT / PDF / Word 文档。", "SYSTEM_ERROR")


def _extract_blocks(md_text: str) -> list[dict]:
    """按 ``=== <id> ===`` 分块；无标记时整篇作为一个块。"""
    parts = _BLOCK_MARKER.split(md_text)
    blocks: list[dict] = []
    i = 1
    while i + 1 < len(parts):
        block_id = parts[i].strip()
        raw = parts[i + 1].strip()
        if raw:
            blocks.append({"block_id": block_id, "raw": raw})
        i += 2
    if not blocks and md_text.strip():
        blocks.append({"block_id": "0", "raw": md_text.strip()})
    return blocks


def _parse_markdown_pipe_table(block: str) -> pd.DataFrame | None:
    """解析 Markdown 管道表格；非表格返回 None。"""
    lines = [ln.strip() for ln in block.splitlines() if ln.strip().startswith("|")]
    if len(lines) < 2:
        return None

    rows: list[list[str]] = []
    for ln in lines:
        # 只去掉首尾各一个边界 |，保留行首空单元格（未命名的行号/索引列）
        s = ln.strip()
        if s.startswith("|"):
            s = s[1:]
        if s.endswith("|"):
            s = s[:-1]
        rows.append([c.strip() for c in s.split("|")])

    header: list[str] | None = None
    data_rows: list[list[str]] = []
    for r in rows:
        if all(re.fullmatch(r":?-+:?", c) for c in r if c):
            continue  # 分隔行（--- / :---: 等）
        if header is None:
            header = r
        else:
            data_rows.append(r)

    if header is None or not data_rows:
        return None
    n = len(header)
    data_rows = [(r + [""] * n)[:n] for r in data_rows]  # 对齐列数
    return pd.DataFrame(data_rows, columns=header)


_HTML_ROW_RE = re.compile(r"<tr\b.*?</tr>", flags=re.S | re.I)
_HTML_CELL_RE = re.compile(r"<t[hd]\b[^>]*>(.*?)</t[hd]>", flags=re.S | re.I)
_HTML_TAG_RE = re.compile(r"<[^>]+>")


def _parse_html_table_fallback(html: str) -> pd.DataFrame | None:
    """无 lxml/html5lib 时的正则回退：提取 <tr>/<th>/<td> 组成 DataFrame。"""
    rows: list[list[str]] = []
    for tr in _HTML_ROW_RE.findall(html):
        cells = [_HTML_TAG_RE.sub("", c).strip() for c in _HTML_CELL_RE.findall(tr)]
        if cells:
            rows.append(cells)
    if len(rows) < 2:
        return None
    header, data_rows = rows[0], rows[1:]
    n = len(header)
    data_rows = [(r + [""] * n)[:n] for r in data_rows]
    return pd.DataFrame(data_rows, columns=header)


def _parse_html_table(block: str) -> pd.DataFrame | None:
    """解析 HTML ``<table>``；优先 pd.read_html（需 lxml 等），失败用正则回退。"""
    m = _HTML_TABLE_RE.search(block)
    if not m:
        return None
    html = m.group(0)
    try:
        dfs = pd.read_html(io.StringIO(html))
    except Exception:
        return _parse_html_table_fallback(html)
    return dfs[0] if dfs else _parse_html_table_fallback(html)


def _ensure_schema(engine) -> None:
    """确保 docs.db 的基础表存在（幂等）。"""
    # 先探明是否真需要迁移：旧 docs.db 的 doc_blocks 缺 columns 列（存原始表头 JSON）。
    # 用表结构判断，而不是盲 ALTER + `except Exception: pass` —— 后者把「列已存在」
    # 这种完全正常的情况与真正的 DDL 失败混成一类，且每次调用都白抛一次异常。
    insp = inspect(engine)
    needs_columns = "doc_blocks" in insp.get_table_names() and "columns" not in {
        c["name"] for c in insp.get_columns("doc_blocks")
    }
    with engine.begin() as conn:
        conn.execute(
            text(
                "CREATE TABLE IF NOT EXISTS doc_meta ("
                "doc_id TEXT PRIMARY KEY, filename TEXT, title TEXT, kind TEXT, created_at TEXT)"
            )
        )
        conn.execute(
            text(
                "CREATE TABLE IF NOT EXISTS doc_blocks ("
                "doc_id TEXT, block_index INTEGER, block_id TEXT, "
                "kind TEXT, content TEXT, table_name TEXT, columns TEXT)"
            )
        )
        if needs_columns:
            conn.execute(text("ALTER TABLE doc_blocks ADD COLUMN columns TEXT"))


def ensure_docs_schema() -> None:
    """在启动期初始化 docs.db 的基础表（幂等，供 ``main.lifespan`` 调用）。

    此前**只有** ``import_document`` 会建表，于是「docs.db 文件已存在但里面没有表」的环境
    （导入中途被打断、工作区文件被外部工具创建/清空等）会让所有假定 ``doc_meta`` /
    ``doc_blocks`` 存在的读取路径失效 —— 启动期的孤儿向量对账会直接以
    ``no such table: doc_meta`` 抛出并打出一整页 traceback。

    另：全新安装时 ``data/`` 目录可能还不存在，SQLite 会以「unable to open database file」
    失败，故这里先建目录（``import_document`` 里那句 mkdir 只覆盖上传路径）。
    """
    settings.docs_db_path.parent.mkdir(parents=True, exist_ok=True)
    _ensure_schema(_docs_engine())


def _existing_doc_tables(engine, doc_id: str) -> list[str]:
    prefix = f"doc_{doc_id}_"
    return [t for t in inspect(engine).get_table_names() if t.startswith(prefix)]


def import_document(file: UploadFile) -> str:
    """解析上传文档，写入 docs.db，返回 ``doc:<doc_id>``。"""
    filename = file.filename or ""
    ext = Path(filename).suffix.lower()
    if ext not in _ALLOWED_DOC_EXTS:
        raise AppError("仅支持 Markdown / TXT / PDF / Word（.md/.txt/.pdf/.docx）文档。", "SYSTEM_ERROR")

    content = read_upload_limited(file)
    if not content:
        raise AppError("文件内容为空。", "SYSTEM_ERROR")

    text_content = _extract_text(filename, content)
    blocks = _extract_blocks(text_content)

    doc_id = uuid.uuid4().hex[:12]
    records: list[dict] = []
    tables_to_write: list[tuple[str, pd.DataFrame]] = []
    table_count = 0
    text_count = 0

    for idx, blk in enumerate(blocks):
        raw = blk["raw"]
        df = _parse_markdown_pipe_table(raw)
        if df is None:
            df = _parse_html_table(raw)
        if df is not None and not df.empty:
            df = df.copy()
            orig_columns = list(df.columns)  # 原始表头（清洗前，如数字表头 0/1）
            df.columns = _sanitize_columns(orig_columns)
            table_name = f"doc_{doc_id}_{idx}"
            tables_to_write.append((table_name, df))
            sample = df.head(5).to_string(index=False)
            records.append({
                "doc_id": doc_id,
                "block_index": idx,
                "block_id": blk["block_id"],
                "kind": "table",
                "content": f"表格列名：{orig_columns}；样本行：\n{sample}",
                "table_name": table_name,
                "columns": json.dumps(orig_columns, ensure_ascii=False),
            })
            table_count += 1
        else:
            chunks = chunk_text(raw, settings.doc_chunk_size, settings.doc_chunk_overlap) or [raw]
            for chunk in chunks:
                records.append({
                    "doc_id": doc_id,
                    "block_index": idx,
                    "block_id": blk["block_id"],
                    "kind": "text",
                    "content": chunk,
                    "table_name": None,
                    "columns": None,
                })
                text_count += 1

    if not records:
        raise AppError("文档中没有可解析的内容。", "SYSTEM_ERROR")

    meta_kind = "table" if table_count and not text_count else ("text" if not table_count else "mixed")

    settings.docs_db_path.parent.mkdir(parents=True, exist_ok=True)
    engine = _docs_engine()
    _ensure_schema(engine)

    with engine.begin() as conn:
        conn.execute(text("DELETE FROM doc_meta WHERE doc_id = :d"), {"d": doc_id})
        conn.execute(text("DELETE FROM doc_blocks WHERE doc_id = :d"), {"d": doc_id})
        for old in _existing_doc_tables(engine, doc_id):
            conn.execute(text(f"DROP TABLE IF EXISTS {engine.dialect.identifier_preparer.quote(old)}"))
        conn.execute(
            text(
                "INSERT INTO doc_meta (doc_id, filename, title, kind, created_at) "
                "VALUES (:doc_id, :filename, :title, :kind, :created_at)"
            ),
            {
                "doc_id": doc_id,
                "filename": filename,
                "title": Path(filename).stem,
                "kind": meta_kind,
                "created_at": _now_iso(),
            },
        )
        for r in records:
            conn.execute(
                text(
                    "INSERT INTO doc_blocks "
                    "(doc_id, block_index, block_id, kind, content, table_name, columns) "
                    "VALUES (:doc_id, :block_index, :block_id, :kind, :content, :table_name, :columns)"
                ),
                r,
            )
        for table_name, df in tables_to_write:
            df.to_sql(table_name, conn, if_exists="replace", index=False)

    # 向量索引：把分块文本（含表格块的「列名+样本行」文本）写入 Milvus。
    # 这一步在 SQLite 提交**之后**，所以失败时必须**补偿删除**已写入的记录：
    # 否则留下一个「数据源列表里看得到、表能被 SQL 查到、但 RAG 检索永远为空」的半成品文档，
    # 而用户在界面上只能看到一个「导入失败」的报错，无从知晓它其实已经半落地。
    try:
        vector_index(
            settings.milvus_doc_collection,
            [r["content"] for r in records],
            [
                {
                    "doc_id": doc_id,
                    "block_id": r["block_id"],
                    "kind": r["kind"],
                    "block_index": r["block_index"],
                }
                for r in records
            ],
            drop_old=False,
        )
    except Exception as e:
        _purge_document(engine, doc_id)
        logger.warning("文档向量化失败，已回滚该文档的库内记录：%s", filename, exc_info=True)
        raise AppError(
            "文档向量化失败，未导入任何内容。请确认向量库（Milvus）与 embedding 服务可用后重试。",
            "SYSTEM_ERROR",
        ) from e

    return f"doc:{doc_id}"


def _docs_engine():
    """docs.db 的引擎（**复用** engine 模块的缓存）。

    早前这里每次都 ``create_engine``：``list_documents`` / ``get_doc_meta`` 等被
    ``GET /api/datasources``、``GET /api/schema`` 反复调用，等于**每个列表请求都新建
    并丢弃一个连接池**（一次 schema 咨询可建 50+ 个）。同一个 SQLite 文件本就有
    ``engine.py`` 里键为 ``"docs"`` 的缓存引擎，两套连接池并存毫无必要。

    复用同一个键还让 ``invalidate_engine("docs")``（删除文档时调用）能同时生效。
    """
    return _get_or_create_engine("docs", settings.docs_db_url, _engine_kwargs("sqlite"))


def get_doc_meta(doc_id: str) -> dict | None:
    """返回文档元信息；不存在返回 None。"""
    if not settings.docs_db_path.exists():
        return None
    try:
        with _docs_engine().connect() as conn:
            row = conn.execute(
                text("SELECT doc_id, filename, title, kind, created_at FROM doc_meta WHERE doc_id = :d"),
                {"d": doc_id},
            ).mappings().first()
    except Exception:
        return None
    return dict(row) if row else None


def get_doc_tables(doc_id: str) -> list[str]:
    """返回该文档提取出的 SQLite 表名列表（无表格为空）。

    与同族读取函数（``list_documents`` / ``get_doc_meta``）保持同一降级口径：
    docs.db 只有文件没有表（旧环境、导入被打断）时返回空，而不是把
    ``no such table: doc_blocks`` 抛到请求层变成 500。
    """
    if not settings.docs_db_path.exists():
        return []
    try:
        with _docs_engine().connect() as conn:
            rows = conn.execute(
                text(
                    "SELECT DISTINCT table_name FROM doc_blocks "
                    "WHERE doc_id = :d AND table_name IS NOT NULL"
                ),
                {"d": doc_id},
            ).scalars().all()
    except Exception:
        logger.warning("读取文档 %s 的表格列表失败，按无表格处理。", doc_id, exc_info=True)
        return []
    return [r for r in rows if r]


def get_doc_table_by_block(doc_id: str, block_id: str) -> str | None:
    """返回某文档指定块（block_id）提取出的表名；无对应表格返回 None。"""
    if not settings.docs_db_path.exists():
        return None
    try:
        with _docs_engine().connect() as conn:
            row = conn.execute(
                text(
                    "SELECT table_name FROM doc_blocks "
                    "WHERE doc_id = :d AND block_id = :b AND table_name IS NOT NULL LIMIT 1"
                ),
                {"d": doc_id, "b": block_id},
            ).scalar()
    except Exception:
        logger.warning("按块 id 定位文档 %s 的表格失败，已跳过。", doc_id, exc_info=True)
        return None
    return row or None


def get_doc_table_original_columns(table_name: str) -> list[str] | None:
    """返回某文档表**清洗前的原始表头**（如数字表头 ``0/1`` 被清洗成 ``t_0/t_1``）；无记录返回 None。

    供 schema 咨询（``build_schema_overview`` / ``answer_schema_question``）用原始名作答，
    避免把「哪一列叫 0」这类问句在清洗后的列名里查不到。
    """
    if not settings.docs_db_path.exists():
        return None
    try:
        with _docs_engine().connect() as conn:
            raw = conn.execute(
                text("SELECT columns FROM doc_blocks WHERE table_name = :t AND columns IS NOT NULL LIMIT 1"),
                {"t": table_name},
            ).scalar()
    except Exception:
        return None  # 旧 docs.db 尚未迁移（无 columns 列）等，静默降级
    if not raw:
        return None
    try:
        cols = json.loads(raw)
    except (TypeError, ValueError):
        return None
    return cols if isinstance(cols, list) else None


def list_documents() -> list[dict]:
    """列出全部文档数据源（供 /api/datasources 聚合）。

    每个文档只产出一个文档级条目（id ``doc:<doc_id>``，name 为文档名）。
    选中后由程序自动读取该文档提取出的全部表格（查询/schema 咨询均如此），
    不再向用户暴露按块表格条目。
    """
    if not settings.docs_db_path.exists():
        return []
    engine = _docs_engine()
    try:
        with engine.connect() as conn:
            rows = conn.execute(
                text("SELECT doc_id, filename, title, kind FROM doc_meta ORDER BY created_at DESC")
            ).mappings().all()
    except Exception:
        return []
    result: list[dict] = []
    for r in rows:
        result.append({
            "id": f"doc:{r['doc_id']}",
            "name": r["title"] or r["filename"] or r["doc_id"],
            "type": "doc",
        })
    return result


def _doc_vector_hits(
    doc_id: str, question: str, top_k: int, with_score: bool
) -> list[dict]:
    """该文档范围内的原始向量命中（按 ``doc_id`` 过滤，绝不跨文档）。"""
    return vector_search(
        settings.milvus_doc_collection,
        question,
        top_k=top_k,
        expr=f'doc_id == "{_safe_doc_id(doc_id)}"',
        with_score=with_score,
    )


def _attach_doc_meta(doc_id: str, hits: list[dict]) -> list[dict]:
    """给向量命中补对外字段；``title`` 取自 doc_meta 而非 Milvus 元数据，历史文档亦可用。

    表格块额外补 ``table_name`` 与 ``row_count``：Milvus 里表格块只索引
    「列名 + 前 5 行样本」，而 docs.db 里存的是**全量数据**。调用方（RAG 回答）需要
    知道「这是摘录、全表其实有 N 行」，否则模型会把前 5 行当全量算聚合，
    一本正经地报出一个假统计量。``row_count`` 查不到时为 None，由调用方措辞兜底。
    """
    meta = get_doc_meta(doc_id) or {}
    title = meta.get("title") or meta.get("filename") or doc_id
    out: list[dict] = []
    for h in hits:
        kind = h.get("kind", "text")
        item = {
            "title": title,
            "block_id": h.get("block_id", ""),
            "block_index": h.get("block_index", 0),
            "kind": kind,
            "content": h["content"],
        }
        if kind == "table":
            # 与块定位共用同一套「块 → 表」映射（get_doc_table_by_block），避免两处逻辑漂移
            table_name = get_doc_table_by_block(doc_id, str(h.get("block_id") or ""))
            if table_name:
                item["table_name"] = table_name
                item["row_count"] = get_doc_table_row_count(table_name)
        if "score" in h:
            item["score"] = h["score"]
        out.append(item)
    return out


def get_doc_table_row_count(table_name: str) -> int | None:
    """返回文档提取表的**总行数**（对照 Milvus 里只有前 5 行样本）；失败返回 None。

    只用于向模型交代「摘录了几行、全表多少行」，不作为查询结果下发，
    故失败时静默降级（调用方按「未知」措辞即可）。
    """
    if not settings.docs_db_path.exists():
        return None
    engine = _docs_engine()
    quoted = engine.dialect.identifier_preparer.quote(table_name)
    try:
        with engine.connect() as conn:
            return int(conn.execute(text(f"SELECT COUNT(*) FROM {quoted}")).scalar() or 0)
    except Exception:
        logger.warning("读取文档表 %s 行数失败，已跳过。", table_name, exc_info=True)
        return None


def retrieve_doc(doc_id: str, question: str, top_k: int | None = None) -> list[dict]:
    """在文档分块中向量检索与问题最相关的 top-k 片段。

    返回 ``[{'title','block_id','block_index','kind','content'}]``，其中 ``title`` 为文档名
    （从 doc_meta 查得，不依赖 Milvus 元数据，历史文档亦可用）；表格块另有
    ``table_name`` / ``row_count``（见 ``_attach_doc_meta``）。
    """
    hits = _doc_vector_hits(doc_id, question, top_k or settings.doc_top_k, with_score=False)
    return _attach_doc_meta(doc_id, hits)


def retrieve_doc_scored(
    doc_id: str,
    question: str,
    top_k: int | None = None,
    max_distance: float | None = None,
    margin: float | None = None,
) -> list[dict]:
    """带阈值过滤的文档检索，用于**表格定位**（返回项额外带 ``score``）。

    区分于 ``retrieve_doc``：那个保召回（宁可多给几段让模型自己看），这个要把
    「文档里根本没有相关内容」与「有相关内容」分开——定位错表比不定位更糟，
    因为收窄后的 schema 会让模型**确定地**用错表。

    Milvus 集合的 metric 是 **L2 距离，越小越相似**，故 ``max_distance`` 是上界：

    - 最优命中都超过它 → 返回空，交由调用方回退全量 schema；
    - 只保留 ``score <= best + margin`` 的命中：内容逐字相同的副本块分数完全一致，
      这条能把它们**成组保留**（漏掉副本无妨，但漏掉「同分的另一种表」就丢了候选），
      同时挡掉次优但不相关的块。
    """
    hits = _doc_vector_hits(doc_id, question, top_k or settings.doc_top_k, with_score=True)
    if not hits:
        return []
    limit = settings.doc_table_max_distance if max_distance is None else max_distance
    best = min(float(h["score"]) for h in hits)
    if best > limit:
        logger.info(
            "文档 %s 的表格定位：最优命中 %.4f 超过阈值 %.2f，判为未定位。",
            doc_id,
            best,
            limit,
        )
        return []
    keep = settings.doc_table_score_margin if margin is None else margin
    return _attach_doc_meta(doc_id, [h for h in hits if float(h["score"]) <= best + keep])


def _purge_document(engine, doc_id: str) -> None:
    """删除某文档在 docs.db 中的全部记录（元信息 / 分块 / 提取出的表）。"""
    old_tables = _existing_doc_tables(engine, doc_id)  # 事务外枚举，避免持有写锁时另开连接 inspect
    with engine.begin() as conn:
        conn.execute(text("DELETE FROM doc_meta WHERE doc_id = :d"), {"d": doc_id})
        conn.execute(text("DELETE FROM doc_blocks WHERE doc_id = :d"), {"d": doc_id})
        for old in old_tables:
            conn.execute(text(f"DROP TABLE IF EXISTS {engine.dialect.identifier_preparer.quote(old)}"))


def delete_document(doc_id: str) -> bool:
    """删除文档数据源：清理 docs.db 的元信息/分块/表格 + Milvus 向量；返回是否删除。"""
    safe_id = _safe_doc_id(doc_id)
    if not settings.docs_db_path.exists() or get_doc_meta(safe_id) is None:
        return False
    engine = _docs_engine()
    _purge_document(engine, safe_id)
    # 库内记录已删，向量若删不掉就成了「查不到、也管不了」的孤儿。此处只记日志：
    # 残留向量永远匹配不到任何已列出的 doc_id（查询恒带 doc_id 过滤），不影响正确性，
    # 但需要留痕以便运维排查；启动时的 reconcile_orphan_vectors() 会兜底清理。
    if not vector_delete(settings.milvus_doc_collection, f'doc_id == "{safe_id}"'):
        logger.error("文档 %s 已删除，但其 Milvus 向量清理失败，已留作孤儿待对账清理。", safe_id)
    invalidate_engine("docs")
    return True


def reconcile_orphan_vectors() -> int:
    """对账：删除 Milvus 中已无 ``doc_meta`` 记录的文档向量；返回清理的 doc_id 数。

    兜底 delete_document 里向量清理失败的残留（Milvus 短暂不可用、进程被重启打断等）。
    纯尽力而为：库或 Milvus 任一不可用都直接跳过，绝不影响启动。
    """
    if not settings.docs_db_path.exists():
        return 0
    engine = _docs_engine()
    # 「文件在、表不在」是尚未导入过任何文档的正常状态，不是异常：直接跳过，
    # 不要用 warning + 整页 traceback 惊动使用者（启动期日志里尤其刺眼）。
    if not inspect(engine).has_table("doc_meta"):
        logger.info("对账孤儿向量：docs.db 尚无 doc_meta 表（未导入过文档），跳过。")
        return 0
    try:
        with engine.connect() as conn:
            known = {
                str(r) for r in conn.execute(text("SELECT doc_id FROM doc_meta")).scalars().all() if r
            }
    except Exception:
        logger.warning("对账孤儿向量：读取 doc_meta 失败，已跳过。", exc_info=True)
        return 0

    orphans = [
        d for d in list_field_values(settings.milvus_doc_collection, "doc_id")
        if _DOC_ID_RE.fullmatch(d) and d not in known
    ]
    for orphan in orphans:
        vector_delete(settings.milvus_doc_collection, f'doc_id == "{orphan}"')
    if orphans:
        logger.info("对账清理了 %d 份孤儿文档向量：%s", len(orphans), ", ".join(orphans))
    return len(orphans)
